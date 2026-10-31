"""
Meeting Exports Service & Views (Section 43)
Generates PDF, DOCX, TXT, SRT, and JSON exports with complete meeting intelligence and full transcript.
"""

import io
import json
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from rest_framework.permissions import AllowAny

from services.storage.mongo import get_collection

def format_timestamp_srt(seconds: float) -> str:
    total_ms = int(seconds * 1000)
    hours = total_ms // 3600000
    rem = total_ms % 3600000
    minutes = rem // 60000
    rem = rem % 60000
    sec = rem // 1000
    ms = rem % 1000
    return f"{hours:02d}:{minutes:02d}:{sec:02d},{ms:03d}"

class ExportPDFView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, meeting_id):
        meetings_col = get_collection("meetings")
        meeting = meetings_col.find_one({"_id": meeting_id})
        if not meeting:
            return Response({"error": "Meeting not found."}, status=status.HTTP_404_NOT_FOUND)

        analysis_col = get_collection("meeting_analysis")
        analysis = analysis_col.find_one({"meetingId": meeting_id}) or {}

        speakers_col = get_collection("speakers")
        speakers = speakers_col.find({"meetingId": meeting_id})

        segments_col = get_collection("transcript_segments")
        segments = segments_col.find({"meetingId": meeting_id}, sort=[("startTime", 1)])

        # Generate PDF using ReportLab
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        story = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor("#1e293b"),
            spaceAfter=10
        )
        heading_style = ParagraphStyle(
            "SectionHeading",
            parent=styles["Heading2"],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#4f46e5"),
            spaceBefore=14,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            "NormalBody",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#334155")
        )

        # Header
        story.append(Paragraph(meeting.get("title", "Meeting Report"), title_style))
        story.append(Paragraph(f"<b>Date:</b> {meeting.get('createdAt', '')} | <b>Duration:</b> {meeting.get('duration', 0)//60} minutes | <b>Status:</b> {meeting.get('status', '').upper()}", body_style))
        story.append(Spacer(1, 10))

        # Participants
        speaker_names = [f"{s.get('name', 'Speaker')} ({s.get('role', 'Participant')})" for s in speakers]
        story.append(Paragraph("<b>Participants:</b> " + (", ".join(speaker_names) if speaker_names else "Team"), body_style))
        story.append(Spacer(1, 12))

        # Executive Summary
        story.append(Paragraph("Executive Summary", heading_style))
        story.append(Paragraph(analysis.get("summary", "No summary available."), body_style))
        story.append(Spacer(1, 12))

        # Key Discussion Points
        story.append(Paragraph("Key Discussion Points", heading_style))
        for kp in analysis.get("keyPoints", []):
            story.append(Paragraph(f"• {kp}", body_style))
        story.append(Spacer(1, 12))

        # Decisions & Final Decision
        story.append(Paragraph("Decisions", heading_style))
        for d in analysis.get("decisions", []):
            story.append(Paragraph(f"• {d}", body_style))

        if analysis.get("finalDecision"):
            story.append(Spacer(1, 6))
            story.append(Paragraph(f"<b>Final Decision:</b> {analysis.get('finalDecision')}", body_style))
        story.append(Spacer(1, 12))

        # Action Items Table
        action_items = analysis.get("actionItems", [])
        if action_items:
            story.append(Paragraph("Action Items", heading_style))
            table_data = [["Task", "Assignee", "Deadline", "Status"]]
            for ai in action_items:
                table_data.append([
                    Paragraph(ai.get("task", ""), body_style),
                    Paragraph(ai.get("assignee", "Unassigned"), body_style),
                    Paragraph(ai.get("deadline", "TBD"), body_style),
                    Paragraph(ai.get("status", "pending").capitalize(), body_style)
                ])
            t = Table(table_data, colWidths=[240, 90, 90, 80])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ]))
            story.append(t)
            story.append(Spacer(1, 14))

        # Full Transcript
        if segments:
            story.append(Paragraph("Full Transcript", heading_style))
            for seg in segments:
                t_sec = int(seg.get("startTime", 0))
                time_str = f"{t_sec//60:02d}:{t_sec%60:02d}"
                speaker = seg.get("speakerName") or seg.get("speakerLabel") or "Speaker"
                story.append(Paragraph(f"<b>[{time_str}] {speaker}:</b> {seg.get('text', '')}", body_style))
                story.append(Spacer(1, 4))

        doc.build(story)
        buffer.seek(0)

        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        filename = f"meeting_{meeting_id}_report.pdf"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class ExportDocxView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, meeting_id):
        meetings_col = get_collection("meetings")
        meeting = meetings_col.find_one({"_id": meeting_id})
        if not meeting:
            return Response({"error": "Meeting not found."}, status=status.HTTP_404_NOT_FOUND)

        analysis_col = get_collection("meeting_analysis")
        analysis = analysis_col.find_one({"meetingId": meeting_id}) or {}

        segments_col = get_collection("transcript_segments")
        segments = segments_col.find({"meetingId": meeting_id}, sort=[("startTime", 1)])

        from docx import Document
        doc = Document()
        doc.add_heading(meeting.get("title", "Meeting Report"), level=0)
        doc.add_paragraph(f"Date: {meeting.get('createdAt')} | Duration: {meeting.get('duration', 0)//60} mins")

        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(analysis.get("summary", ""))

        doc.add_heading("Key Discussion Points", level=1)
        for p in analysis.get("keyPoints", []):
            doc.add_paragraph(p, style='List Bullet')

        doc.add_heading("Decisions & Final Decision", level=1)
        for d in analysis.get("decisions", []):
            doc.add_paragraph(d, style='List Bullet')
        if analysis.get("finalDecision"):
            doc.add_paragraph(f"Final Decision: {analysis.get('finalDecision')}")

        doc.add_heading("Action Items", level=1)
        table = doc.add_table(rows=1, cols=4)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Task'
        hdr_cells[1].text = 'Assignee'
        hdr_cells[2].text = 'Deadline'
        hdr_cells[3].text = 'Status'
        for ai in analysis.get("actionItems", []):
            row_cells = table.add_row().cells
            row_cells[0].text = ai.get("task", "")
            row_cells[1].text = ai.get("assignee", "")
            row_cells[2].text = ai.get("deadline", "")
            row_cells[3].text = ai.get("status", "")

        doc.add_heading("Transcript", level=1)
        for s in segments:
            t_sec = int(s.get("startTime", 0))
            time_str = f"{t_sec//60:02d}:{t_sec%60:02d}"
            spk = s.get("speakerName") or s.get("speakerLabel") or "Speaker"
            doc.add_paragraph(f"[{time_str}] {spk}: {s.get('text', '')}")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        response["Content-Disposition"] = f'attachment; filename="meeting_{meeting_id}.docx"'
        return response


class ExportTxtView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, meeting_id):
        meetings_col = get_collection("meetings")
        meeting = meetings_col.find_one({"_id": meeting_id})
        if not meeting:
            return Response({"error": "Meeting not found."}, status=status.HTTP_404_NOT_FOUND)

        analysis_col = get_collection("meeting_analysis")
        analysis = analysis_col.find_one({"meetingId": meeting_id}) or {}

        segments_col = get_collection("transcript_segments")
        segments = segments_col.find({"meetingId": meeting_id}, sort=[("startTime", 1)])

        lines = [
            f"TITLE: {meeting.get('title')}",
            f"DATE: {meeting.get('createdAt')}",
            f"STATUS: {meeting.get('status')}",
            "=" * 60,
            "\nEXECUTIVE SUMMARY:",
            analysis.get("summary", ""),
            "\nKEY POINTS:",
        ]
        for kp in analysis.get("keyPoints", []):
            lines.append(f"- {kp}")

        lines.extend([
            "\nDECISIONS:",
        ])
        for d in analysis.get("decisions", []):
            lines.append(f"- {d}")
        lines.append(f"\nFINAL DECISION:\n{analysis.get('finalDecision', '')}")

        lines.extend(["\nACTION ITEMS:"])
        for ai in analysis.get("actionItems", []):
            lines.append(f"- [{ai.get('status')}] {ai.get('task')} (Assignee: {ai.get('assignee')}, Deadline: {ai.get('deadline')})")

        lines.extend(["\nTRANSCRIPT:", "=" * 60])
        for s in segments:
            t = int(s.get("startTime", 0))
            lines.append(f"[{t//60:02d}:{t%60:02d}] {s.get('speakerName', 'Speaker')}: {s.get('text')}")

        content = "\n".join(lines)
        response = HttpResponse(content, content_type="text/plain; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="meeting_{meeting_id}.txt"'
        return response


class ExportSRTView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, meeting_id):
        segments_col = get_collection("transcript_segments")
        segments = segments_col.find({"meetingId": meeting_id}, sort=[("startTime", 1)])

        srt_blocks = []
        for idx, s in enumerate(segments, 1):
            start = format_timestamp_srt(s.get("startTime", 0.0))
            end = format_timestamp_srt(s.get("endTime", s.get("startTime", 0.0) + 3.0))
            spk = s.get("speakerName") or s.get("speakerLabel") or ""
            prefix = f"{spk}: " if spk else ""
            text = f"{prefix}{s.get('text', '')}"
            srt_blocks.append(f"{idx}\n{start} --> {end}\n{text}\n")

        content = "\n".join(srt_blocks)
        response = HttpResponse(content, content_type="text/plain; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="meeting_{meeting_id}.srt"'
        return response


class ExportJSONView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, meeting_id):
        meetings_col = get_collection("meetings")
        meeting = meetings_col.find_one({"_id": meeting_id})
        if not meeting:
            return Response({"error": "Meeting not found."}, status=status.HTTP_404_NOT_FOUND)

        analysis_col = get_collection("meeting_analysis")
        analysis = analysis_col.find_one({"meetingId": meeting_id}) or {}

        speakers_col = get_collection("speakers")
        speakers = speakers_col.find({"meetingId": meeting_id})

        segments_col = get_collection("transcript_segments")
        segments = segments_col.find({"meetingId": meeting_id}, sort=[("startTime", 1)])

        data = {
            "meeting": meeting,
            "speakers": speakers,
            "analysis": analysis,
            "transcript": segments
        }

        response = HttpResponse(json.dumps(data, indent=2, default=str), content_type="application/json")
        response["Content-Disposition"] = f'attachment; filename="meeting_{meeting_id}.json"'
        return response
